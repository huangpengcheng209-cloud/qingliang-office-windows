"""Exercise real Qt/CEF PPTX editing, native save/reopen, slideshow and PDF export.

Input uses Chromium events. Multimedia, object animations, native IME and installers remain separate.
"""
import argparse
import json
import os
import shutil
import zipfile
from pathlib import Path

from run_editor_tests import Acceptance, ROOT, owned_dialog, sha


def create_presentation(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, title in enumerate(['氢量演示验收 Presentation', '第二页 · 流程与数据']):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(.8), Inches(11), Inches(.9))
        box.name, box.text = '验收标题', title
        run = box.text_frame.paragraphs[0].runs[0]
        run.font.name, run.font.size = 'Microsoft YaHei', Pt(30)
        run.font.color.rgb = RGBColor.from_string('163358')
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2.2), Inches(4), Inches(2))
        shape.name, shape.text = '流程形状', '保留形状与内容'
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string('2563EB')
        if index == 1:
            table = slide.shapes.add_table(3, 2, Inches(6), Inches(2.2), Inches(5), Inches(2)).table
            for row, values in enumerate([['阶段', '状态'], ['编辑', '可用'], ['验收', '进行中']]):
                for column, value in enumerate(values): table.cell(row, column).text = value
        slide.notes_slide.notes_text_frame.text = '演讲者备注 ' + str(index + 1)
        transition = OxmlElement('p:transition')
        transition.set('spd', 'med')
        transition.append(OxmlElement('p:fade'))
        slide._element.append(transition)
    deck.save(path)


class PresentationAcceptance(Acceptance):
    def __init__(self, args):
        super().__init__(args)
        self.source = self.evidence / 'fixtures/source.pptx'
        self.edited = self.evidence / 'fixtures/演示 工作副本.pptx'
        self.created = self.evidence / 'fixtures/新建演示.PPTX'
        self.marker = '真实演示编辑 Native slide 2026'
        self.new_marker = '新建演示验收 New presentation'

    def fixtures(self):
        marker = self.evidence / '.qingliang-presentation-acceptance'
        if self.evidence.exists() and any(self.evidence.iterdir()) and not marker.exists():
            raise ValueError('Select an empty evidence directory or a marked presentation acceptance directory')
        self.evidence.mkdir(parents=True, exist_ok=True)
        marker.write_text('Only controlled presentation acceptance fixtures and evidence.\n', 'utf-8')
        self.source.parent.mkdir(exist_ok=True)
        result = self.evidence / 'presentation-test-result.json'
        known = json.loads(result.read_text('utf-8')).get('owned_fixtures', {}) if result.exists() else {}
        paths = (self.source, self.edited, self.created, self.pdf)
        for path in paths:
            if path.exists() and sha(path) != known.get(path.relative_to(self.evidence).as_posix()):
                raise ValueError('Unknown fixture edit preserved: ' + path.name)
        for path in (self.edited, self.created, self.pdf):
            if path.exists(): path.unlink()
        if not self.source.exists(): create_presentation(self.source)
        shutil.copy2(self.source, self.edited)
        self.original_hash = sha(self.source)
        self.fixtures_ready = True

    def connect(self, playwright):
        browser = playwright.chromium.connect_over_cdp('http://127.0.0.1:' + str(self.port))
        page = next(p for c in browser.contexts for p in c.pages if '/apps/api/documents/' in p.url)
        page.set_default_timeout(7000)
        page.on('pageerror', lambda error: self.errors.append(str(error)))
        for attempt in range(120):
            frame = next((f for f in page.frames if '/presentationeditor/main/' in f.url), None)
            if frame and frame.evaluate('(()=>{try{return editor.getCountPages()>0 && !!document.querySelector(".ql-quickbar")}catch{return false}})()'):
                frame.locator('.ql-quickbar').wait_for(state='visible')
                return browser, page, frame
            page.wait_for_timeout(250)
        raise AssertionError('The real presentation editor did not finish rendering')

    def shapes(self, frame):
        return frame.evaluate('AscBuilder.Slide.Api.GetPresentation().GetCurrentSlide().GetAllShapes().map(s=>({name:s.GetName(),text:s.GetDocContent()?.GetText()}))')

    def wait_saved(self, page, frame, path, predicate):
        from pptx import Presentation
        from pptx.exc import PackageNotFoundError
        previous, stable = None, 0
        for attempt in range(120):
            page.wait_for_timeout(150)
            try:
                ready = predicate(Presentation(path))
                current = sha(path)
                stable = stable + 1 if ready and current == previous else 0
                previous = current
                if ready and stable >= 3 and not frame.evaluate('editor.asc_isDocumentCanSave()'): return
            except (OSError, ValueError, KeyError, zipfile.BadZipFile, PackageNotFoundError): pass
        raise AssertionError('The native PPTX save did not reach independently verified stable output')

    def edit_existing(self, playwright):
        from pptx import Presentation
        self.start()
        self.open_argument(self.edited)
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        self.check('existing_pptx_native_source_path', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.edited)
        self.check('existing_pptx_has_two_slides', frame.evaluate('editor.getCountPages()') == 2)
        self.check('initial_title_and_shape_in_real_model', [s['text'].strip() for s in self.shapes(frame)] == ['氢量演示验收 Presentation', '保留形状与内容'])
        # Scoped canvas coordinates in the default isolated preview layout; never global desktop input.
        frame.locator('#id_viewer_overlay').dblclick(position={'x': 220, 'y': 110})
        page.keyboard.press('Control+A')
        self.check('selects_only_slide_title_text', frame.evaluate('editor.pluginMethod_GetSelectedText({Numbering:false})').strip() == '氢量演示验收 Presentation')
        page.keyboard.insert_text(self.marker)
        frame.wait_for_function('editor.asc_isDocumentCanSave()')
        self.check('editing_enables_native_save', frame.locator('.ql-quickbar [data-ql-action="save"]').is_enabled())
        page.keyboard.press('Control+Z')
        self.check('undo_restores_original_slide_title', self.shapes(frame)[0]['text'].strip() == '氢量演示验收 Presentation')
        frame.locator('.ql-quickbar [data-ql-action="redo"]').click()
        self.check('redo_restores_edited_slide_title', self.shapes(frame)[0]['text'].strip() == self.marker)
        page.keyboard.press('Control+A')
        frame.locator('.ql-quickbar [data-ql-action="bold"]').click()
        self.check('bold_control_reports_selected_text_state', frame.locator('.ql-quickbar [data-ql-action="bold"]').get_attribute('aria-pressed') == 'true')
        page.keyboard.press('End')
        page.keyboard.insert_text(' 后续输入')
        self.marker += ' 后续输入'
        self.check('formatting_returns_focus_for_continued_typing', self.shapes(frame)[0]['text'].strip() == self.marker)
        page.keyboard.press('Escape')
        frame.locator('#id_viewer_overlay').click(position={'x': 100, 'y': 205})
        page.keyboard.press('ArrowRight')
        page.keyboard.press('ArrowDown')
        frame.locator('.ql-quickbar [data-ql-action="slide"]').click()
        self.check('quick_add_inserts_slide_after_current', frame.evaluate('editor.getCountPages()===3 && editor.getCurrentPage()===1'))
        page.keyboard.press('Control+Z')
        self.check('undo_removes_added_slide', frame.evaluate('editor.getCountPages()') == 2)
        frame.locator('.ql-quickbar [data-ql-action="redo"]').click()
        self.check('redo_restores_added_slide', frame.evaluate('editor.getCountPages()') == 3)
        frame.locator('.ql-quickbar [data-ql-action="save"]').click()
        self.wait_saved(page, frame, self.edited, lambda deck: len(deck.slides) == 3 and deck.slides[0].shapes[0].text == self.marker)
        saved, source = Presentation(self.edited), Presentation(self.source)
        self.check('independent_saved_title_matches_keyboard_edit', saved.slides[0].shapes[0].text == self.marker)
        self.check('independent_saved_title_retains_bold', all(run.font.bold for p in saved.slides[0].shapes[0].text_frame.paragraphs for run in p.runs))
        moved, original = saved.slides[0].shapes[1], source.slides[0].shapes[1]
        self.check('shape_keyboard_move_saved_on_both_axes', moved.left > original.left and moved.top > original.top)
        self.check('shape_text_size_and_fill_preserved', moved.text == original.text and moved.width == original.width and moved.height == original.height and moved.fill.fore_color.rgb == original.fill.fore_color.rgb)
        self.check('existing_table_values_preserved', [[c.text for c in r.cells] for r in saved.slides[2].shapes[2].table.rows] == [['阶段', '状态'], ['编辑', '可用'], ['验收', '进行中']])
        self.check('speaker_notes_preserved', saved.slides[0].notes_slide.notes_text_frame.text == '演讲者备注 1' and saved.slides[2].notes_slide.notes_text_frame.text == '演讲者备注 2')
        # The converter writes a valid mc:Choice/mc:Fallback pair for newer transition attributes.
        self.check('existing_fade_transition_preserved', all(slide._element.xpath('.//p:transition[@spd="med"]/p:fade') for slide in (saved.slides[0], saved.slides[2])))
        self.slideshow(page, frame)
        page.screenshot(path=str(self.evidence / 'edited-presentation.png'))
        self.export_pdf(page, frame)
        self.close()

    def slideshow(self, page, frame):
        self.check('quick_slideshow_control_enabled', frame.locator('.ql-quickbar [data-ql-action="slideshow"]').is_enabled())
        self.check('slideshow_entry_stays_in_fixed_tools_within_viewport', frame.locator('.ql-tools [data-ql-action="slideshow"]').evaluate('(e)=>{const r=e.getBoundingClientRect();return r.left>=0 && r.right<=innerWidth}'))
        current = frame.evaluate('editor.getCurrentPage()')
        frame.locator('.ql-quickbar [data-ql-action="slideshow"]').click()
        frame.locator('#pe-preview').wait_for(state='visible')
        frame.wait_for_function('n=>PE.getController("Viewport").previewPanel.pages.get("current")===n', arg=current)
        self.check('quick_slideshow_starts_at_current_slide', frame.locator('#pe-preview').is_visible())
        page.keyboard.press('ArrowRight')
        frame.wait_for_function('PE.getController("Viewport").previewPanel.pages.get("current")===2')
        self.check('slideshow_keyboard_next_reaches_existing_second_slide', '3' in frame.locator('#preview-label-slides').inner_text())
        page.wait_for_timeout(800)
        frame.locator('#pe-preview').hover(position={'x': 200, 'y': 200})
        page.screenshot(path=str(self.evidence / 'slideshow-last-slide.png'))
        self.check('moving_pointer_reveals_slideshow_controls', frame.locator('#btn-preview-prev').is_visible())
        frame.locator('#btn-preview-prev').click()
        frame.wait_for_function('PE.getController("Viewport").previewPanel.pages.get("current")===1')
        self.check('slideshow_previous_button_changes_slide', True)
        page.keyboard.press('Escape')
        frame.locator('#pe-preview').wait_for(state='hidden')
        self.check('escape_returns_from_slideshow_to_editor', frame.locator('.ql-quickbar').is_visible())
        frame.locator('.ql-quickbar [data-ql-action="commands"]').click()
        frame.locator('.ql-command-search').fill('从头')
        frame.locator('.ql-command-item[data-ql-action="slideshow-first"]').click()
        frame.locator('#pe-preview').wait_for(state='visible')
        frame.wait_for_function('PE.getController("Viewport").previewPanel.pages.get("current")===0')
        self.check('from_beginning_command_starts_first_slide', True)
        page.wait_for_timeout(800)
        page.screenshot(path=str(self.evidence / 'slideshow-first-slide.png'))
        page.keyboard.press('Escape')
        frame.locator('#pe-preview').wait_for(state='hidden')

    def export_pdf(self, page, frame):
        import fitz
        before = sha(self.edited)
        frame.locator('.ql-quickbar [data-ql-action="file"]').click()
        frame.locator('.ql-command-item[data-ql-action="exportpdf"]').click()
        dialog = owned_dialog(self.pid)
        self.check('pptx_pdf_export_opens_owned_native_dialog', dialog.exists())
        dialog.child_window(auto_id='1001', control_type='Edit', visible_only=False).set_edit_text(str(self.pdf))
        dialog.child_window(auto_id='1', control_type='Button', visible_only=False).invoke()
        for attempt in range(120):
            page.wait_for_timeout(150)
            try:
                with fitz.open(self.pdf) as document:
                    if len(document) == 3 and ''.join(self.marker.split()) in ''.join(document[0].get_text().split()):
                        text = ''.join(p.get_text() for p in document)
                        image = document[0].get_pixmap(matrix=fitz.Matrix(1, 1))
                        image.save(self.evidence / 'exported-first-slide.png')
                        break
            except (OSError, RuntimeError, ValueError): pass
        else: raise AssertionError('Native presentation PDF export not independently readable')
        self.check('exported_pdf_contains_edited_and_original_slides', ''.join(self.marker.split()) in ''.join(text.split()) and '流程与数据' in text and '进行中' in text)
        self.check('exported_pdf_has_landscape_slide_geometry', image.width > image.height)
        self.check('pdf_export_preserves_pptx_bytes', sha(self.edited) == before)
        self.check('pdf_export_preserves_native_pptx_context', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.edited)

    def reopen(self, playwright):
        self.start()
        self.open_argument(self.edited)
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        self.check('fresh_process_reopens_three_slides', frame.evaluate('editor.getCountPages()') == 3)
        frame.locator('#id_thumbnails').click(position={'x': 100, 'y': 70})
        self.check('fresh_process_reopens_saved_title', self.shapes(frame)[0]['text'].strip() == self.marker)
        page.screenshot(path=str(self.evidence / 'reopened-presentation.png'))
        self.close()

    def create_new(self, playwright):
        self.start()
        self.open_argument('--new:slide')
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        self.check('native_new_presentation_opens_one_slide', frame.evaluate('editor.getCountPages()') == 1)
        frame.locator('#id_viewer_overlay').dblclick(position={'x': 220, 'y': 205})
        page.keyboard.insert_text(self.new_marker)
        self.check('native_new_slide_accepts_title_input', any(self.new_marker in (s['text'] or '') for s in self.shapes(frame)))
        frame.locator('.ql-quickbar [data-ql-action="file"]').click()
        frame.locator('.ql-command-item[data-ql-action="saveas"]').click()
        dialog = owned_dialog(self.pid)
        self.check('new_presentation_uses_owned_native_save_dialog', dialog.exists())
        dialog.child_window(auto_id='1001', control_type='Edit', visible_only=False).set_edit_text(str(self.created))
        dialog.child_window(auto_id='1', control_type='Button', visible_only=False).invoke()
        self.wait_saved(page, frame, self.created, lambda deck: len(deck.slides) == 1 and any(self.new_marker in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame))
        self.check('uppercase_pptx_filename_has_no_duplicate_extension', self.created.exists() and not Path(str(self.created) + '.pptx').exists())
        self.check('native_new_presentation_changes_to_saved_source_path', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.created)
        self.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', required=True, type=Path)
    parser.add_argument('--qt-dir', required=True, type=Path)
    parser.add_argument('--evidence-dir', required=True, type=Path)
    args = parser.parse_args()
    if os.name != 'nt': parser.error('Windows is required')
    work = args.work_dir.resolve()
    if not (work / '.qingliang-dependency-build').is_file(): raise ValueError('Select the marked stable build directory')
    build = json.loads((work / 'editor-build-report.json').read_text('utf-8'))
    if build['input_manifest_sha256'] != sha(ROOT / 'editor.inputs.lock.json'): raise ValueError('Build the current editor resources first')
    os.environ['TEMP'] = os.environ['TMP'] = str(work / 'temp')
    from playwright.sync_api import sync_playwright
    acceptance = PresentationAcceptance(args)
    report = {'cases': acceptance.cases, 'observed_javascript_errors': acceptance.errors, 'normal_shutdowns': acceptance.closes,
              'input_method': 'Chromium events; native IME and clipboard not tested',
              'exe_sha256': sha(work / 'bin/DesktopEditors.exe'), 'host_dll_sha256': sha(work / 'bin/ascdocumentscore.dll'),
              'editor_build_report_sha256': sha(work / 'editor-build-report.json'), 'runtime_layout_sha256': sha(work / 'editor-runtime-layout.json'),
              'test_source_sha256': sha(Path(__file__)), 'native_session_helper_sha256': sha(ROOT / 'tests/run_editor_tests.py'),
              'preview_launcher_sha256': sha(ROOT / 'tools/launch_qt_preview.py'),
              'object_animation_acceptance': False, 'multimedia_acceptance': False, 'transition_timing_acceptance': False,
              'pdf_form_editing_acceptance': False, 'installer_acceptance': False}
    try:
        acceptance.fixtures()
        with sync_playwright() as playwright:
            acceptance.edit_existing(playwright)
            acceptance.reopen(playwright)
            acceptance.create_new(playwright)
        acceptance.check('untouched_source_pptx_preserved', sha(acceptance.source) == acceptance.original_hash)
        acceptance.check('interactions_have_no_observed_javascript_errors', not acceptance.errors)
    except BaseException as error:
        report['execution_error'] = str(error)
        raise
    finally:
        if acceptance.fixtures_ready:
            report['owned_fixtures'] = {p.relative_to(acceptance.evidence).as_posix(): sha(p) for p in (acceptance.source, acceptance.edited, acceptance.created, acceptance.pdf) if p.exists()}
            report['passed'] = sum(c['passed'] for c in acceptance.cases)
            report['failed'] = max(sum(not c['passed'] for c in acceptance.cases), int('execution_error' in report))
            (acceptance.evidence / 'presentation-test-result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', 'utf-8')
            print(json.dumps({'passed': report['passed'], 'failed': report['failed'], 'installer_acceptance': False}), flush=True)
        elif acceptance.evidence.exists() and (acceptance.evidence / '.qingliang-presentation-acceptance').exists():
            (acceptance.evidence / 'fixture-preflight-error.json').write_text(json.dumps({'execution_error': report.get('execution_error'), 'prior_ownership_report_preserved': True}, indent=2) + '\n', 'utf-8')


if __name__ == '__main__': main()
