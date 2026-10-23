"""Verify a real Qt/CEF XLSX lifecycle and theme-aware editor controls without a native mock.

Uses Chromium input events; native IME/clipboard, chart rendering and exhaustive spreadsheet acceptance are separate.
"""
import argparse
import json
import os
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from run_editor_tests import Acceptance, ROOT, owned_dialog, sha


def create_workbook(path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = '预算明细'
    sheet.merge_cells('A1:E1')
    sheet['A1'] = '氢量办公表格验收 Budget'
    sheet['A1'].font = Font(size=16, bold=True)
    sheet.append([])
    sheet.append(['项目', '数量', '单价', '金额', '日期'])
    for row, values in enumerate([['纸张', 2, 12.5], ['文件夹', 3, 8], ['墨盒', 1, 80]], 4):
        for column, value in enumerate(values, 1): sheet.cell(row, column, value)
        sheet.cell(row, 4, '=B' + str(row) + '*C' + str(row)).number_format = '0.00'
        sheet.cell(row, 5, datetime(2026, 10, row + 19)).number_format = 'yyyy-mm-dd'
    sheet.auto_filter.ref = 'A3:E6'
    sheet.freeze_panes = 'B4'
    for cell in sheet[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='2563EB')
    for column in 'ABCDE': sheet.column_dimensions[column].width = 18
    summary = book.create_sheet('汇总')
    summary.append(['统计', '合计'])
    summary.append(['金额合计', "=SUM('预算明细'!D4:D6)"])
    chart = BarChart()
    chart.title = '采购金额'
    chart.add_data(Reference(sheet, min_col=4, min_row=3, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=4, max_row=6))
    sheet.add_chart(chart, 'G3')
    book.save(path)


def contrast(foreground, background):
    def color(value):
        values = [float(v) for v in re.findall(r'[\d.]+', value)]
        if len(values) not in (3, 4): raise ValueError('Unexpected computed color: ' + value)
        return values[:3], values[3] if len(values) == 4 else 1
    bg, alpha = color(background)
    if alpha != 1: raise ValueError('The tested surface must be opaque')
    fg, alpha = color(foreground)
    fg = [c * alpha + b * (1 - alpha) for c, b in zip(fg, bg)]
    def luminance(rgb):
        channels = [v / 255 / 12.92 if v / 255 <= .04045 else ((v / 255 + .055) / 1.055) ** 2.4 for v in rgb]
        return sum(a * b for a, b in zip(channels, [.2126, .7152, .0722]))
    a, b = sorted([luminance(fg), luminance(bg)])
    return (b + .05) / (a + .05)


class SpreadsheetAcceptance(Acceptance):
    def __init__(self, args):
        super().__init__(args)
        self.source = self.evidence / 'fixtures/source.xlsx'
        self.edited = self.evidence / 'fixtures/预算 工作副本.xlsx'
        self.created = self.evidence / 'fixtures/新建表格.XLSX'
        self.marker = '表格真实编辑 Native spreadsheet'
        self.style_observations = []

    def fixtures(self):
        marker = self.evidence / '.qingliang-spreadsheet-acceptance'
        if self.evidence.exists() and any(self.evidence.iterdir()) and not marker.exists():
            raise ValueError('Select an empty evidence directory or an existing marked spreadsheet acceptance directory')
        self.evidence.mkdir(parents=True, exist_ok=True)
        marker.write_text('Only controlled spreadsheet acceptance fixtures and evidence.\n', 'utf-8')
        self.source.parent.mkdir(exist_ok=True)
        result = self.evidence / 'spreadsheet-test-result.json'
        known = json.loads(result.read_text('utf-8')).get('owned_fixtures', {}) if result.exists() else {}
        for path in (self.source, self.edited, self.created):
            if path.exists():
                if sha(path) != known.get(path.relative_to(self.evidence).as_posix()): raise ValueError('Unknown fixture edit preserved: ' + path.name)
        for path in (self.edited, self.created):
            if path.exists(): path.unlink()
        if not self.source.exists(): create_workbook(self.source)
        shutil.copy2(self.source, self.edited)
        self.original_hash = sha(self.source)
        self.fixtures_ready = True

    def connect(self, playwright):
        browser = playwright.chromium.connect_over_cdp('http://127.0.0.1:' + str(self.port))
        page = next(p for c in browser.contexts for p in c.pages if '/apps/api/documents/' in p.url)
        page.set_default_timeout(7000)
        page.on('pageerror', lambda error: self.errors.append(str(error)))
        frame = None
        for attempt in range(100):
            frame = next((f for f in page.frames if '/spreadsheeteditor/main/' in f.url), None)
            if frame and frame.evaluate("(() => {try {return Asc.editor.asc_getWorksheetsCount() > 0 && !!document.querySelector('#ce-cell-name');} catch {return false;}})()"):
                frame.locator('.ql-quickbar').wait_for(state='visible')
                return browser, page, frame
            page.wait_for_timeout(200)
        raise AssertionError('The real spreadsheet editor did not finish rendering')

    def go(self, frame, address):
        frame.locator('#ce-cell-name').fill(address)
        frame.locator('#ce-cell-name').press('Enter')
        frame.wait_for_function(r'wanted => Asc.editor.asc_getActiveRangeStr().replace(/\$/g, "") === wanted', arg=address)

    def read(self, frame, address):
        self.go(frame, address)
        return frame.evaluate('Asc.editor.asc_getCellInfo().asc_getText()')

    def put(self, frame, address, text):
        self.go(frame, address)
        frame.locator('#ce-cell-content').fill(text)
        frame.locator('#ce-cell-content').press('Enter')

    def wait_saved(self, page, frame, path, predicate):
        from openpyxl import load_workbook
        previous, stable = None, 0
        for attempt in range(120):
            page.wait_for_timeout(150)
            try:
                book = load_workbook(path, data_only=False)
                try: ready = predicate(book)
                finally: book.close()
                current = sha(path)
                stable = stable + 1 if ready and current == previous else 0
                previous = current
                if stable >= 3 and frame.evaluate('Asc.editor.asc_isDocumentCanSave()') is False: return
            except (OSError, ValueError, KeyError, zipfile.BadZipFile): pass
        raise AssertionError('The expected XLSX content did not finish writing to disk')

    def style_and_motion(self, browser, page, frame):
        home = next(p for c in browser.contexts for p in c.pages if p.url.startswith((self.work / 'bin/index.html').as_uri()))
        home.on('pageerror', lambda error: self.errors.append(str(error)))
        for theme in ['theme-dark', 'theme-light']:
            home.locator('a').filter(has_text='设置').click()
            select = home.locator('#box-settings select').filter(has=home.locator('option[value="theme-dark"]'))
            select.select_option(theme)
            home.locator('#sett-btn-apply:visible').click()
            expected = theme == 'theme-dark'
            frame.wait_for_function("dark => { const rgb=getComputedStyle(document.querySelector('.ql-quickbar')).backgroundColor.match(/[\\d.]+/g).map(Number);return (rgb[0]<128)===dark; }", arg=expected)
            observation = frame.locator('.ql-quickbar').evaluate("e=>({background:getComputedStyle(e).backgroundColor,color:getComputedStyle(e.querySelector('[data-ql-action=commands]')).color})")
            observation.update({'theme': theme, 'toolbar_text_contrast': round(contrast(observation['color'], observation['background']), 2)})
            self.style_observations.append(observation)
            self.check(theme + '_toolbar_text_contrast', observation['toolbar_text_contrast'] >= 4.5)
            self.check(theme + '_editor_viewport_has_no_horizontal_overflow', frame.evaluate('document.documentElement.scrollWidth <= innerWidth'))
            frame.locator('.ql-quickbar [data-ql-action="file"]').click()
            panel = frame.locator('.ql-command-panel')
            panel.wait_for(state='visible')
            frame.wait_for_function("!document.querySelector('.ql-command-panel').getAnimations().some(a=>a.playState==='running')")
            colors = panel.evaluate("e=>({background:getComputedStyle(e).backgroundColor,color:getComputedStyle(e.querySelector('strong')).color})")
            self.check(theme + '_command_panel_follows_native_surface', colors['background'] == observation['background'])
            self.check(theme + '_command_text_contrast', contrast(colors['color'], colors['background']) >= 4.5)
            page.screenshot(path=str(self.evidence / (theme + '-commands.png')))
            frame.locator('.ql-command-search').press('ArrowDown')
            self.check(theme + '_arrow_key_focuses_command', frame.evaluate("document.activeElement.matches('.ql-command-item')"))
            page.keyboard.press('Escape')
            self.check(theme + '_escape_restores_opener_focus', frame.locator('[data-ql-action="file"]').evaluate('e=>e===document.activeElement'))
        page.emulate_media(reduced_motion='reduce')
        self.check('reduced_motion_disables_control_transitions', frame.locator('.ql-quickbar [data-ql-action="file"]').evaluate("e=>getComputedStyle(e).transitionDuration==='0s'"))
        frame.locator('.ql-quickbar [data-ql-action="file"]').click()
        self.check('reduced_motion_disables_panel_animation', frame.locator('.ql-command-panel').evaluate("e=>getComputedStyle(e).animationName==='none'"))
        page.keyboard.press('Escape')
        page.emulate_media(reduced_motion='no-preference')
        home.locator('a').filter(has_text='设置').click()
        select.select_option('theme-system')
        home.locator('#sett-btn-apply:visible').click()

    def inspect_saved(self):
        from openpyxl import load_workbook
        book = load_workbook(self.edited, data_only=False)
        try:
            sheet = book['预算明细']
            self.check('saved_workbook_preserves_original_and_added_sheets', book.sheetnames == ['预算明细', self.new_sheet, '汇总'])
            self.check('saved_sort_moves_complete_rows', [[sheet.cell(r, c).value for c in range(1, 4)] for r in range(4, 7)] == [['墨盒', 1, 80], ['文件夹', 3, 8], ['纸张', 5, 12.5]])
            self.check('saved_workbook_preserves_merged_title_and_filter', str(sheet.merged_cells) == 'A1:E1' and sheet.auto_filter.ref == 'A3:E6')
            self.check('saved_workbook_preserves_number_and_date_formats', sheet['D4'].number_format == '0.00' and sheet['E4'].value == datetime(2026, 10, 25) and sheet['E4'].number_format.lower() == 'yyyy-mm-dd')
            self.check('saved_workbook_preserves_chart_structure', len(sheet._charts) == 1 and len(sheet._charts[0].series) == 1)
            self.check('saved_text_and_bold_are_independently_readable', sheet['A8'].value == self.marker and sheet['A8'].font.bold)
            self.check('saved_percentage_has_numeric_value_and_format', sheet['F8'].value == .125 and '%' in sheet['F8'].number_format)
            self.check('saved_formula_and_cross_sheet_reference_are_preserved', sheet['D7'].value == '=SUM(D4:D6)' and book[self.new_sheet]['B1'].value in ["='预算明细'!D7", '=预算明细!D7'])
        finally: book.close()
        book = load_workbook(self.edited, data_only=True)
        try:
            sheet = book['预算明细']
            self.check('independent_cached_values_match_actual_inputs', [sheet.cell(r, 4).value for r in range(4, 8)] == [80, 24, 62.5, 166.5])
            self.check('independent_cross_sheet_cached_values_recalculate', book['汇总']['B2'].value == 166.5 and book[self.new_sheet]['B1'].value == 166.5)
        finally: book.close()

    def edit_existing(self, playwright):
        self.start()
        self.open_argument(self.edited)
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        self.check('actual_native_source_path_matches_working_copy', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.edited)
        self.check('initial_multi_sheet_model_loaded', frame.evaluate('Array.from({length:Asc.editor.asc_getWorksheetsCount()},(_,i)=>Asc.editor.asc_getWorksheetName(i))') == ['预算明细', '汇总'])
        self.check('initial_input_and_formula_loaded', self.read(frame, 'B4') == '2' and self.read(frame, 'D4') == '=B4*C4')
        self.put(frame, 'B4', '5')
        self.put(frame, 'D7', '=SUM(D4:D6)')
        self.put(frame, 'A8', self.marker)
        page.keyboard.press('Control+Z')
        self.check('undo_removes_new_cell_text', not self.read(frame, 'A8'))
        frame.locator('.ql-quickbar [data-ql-action="redo"]').click()
        self.check('redo_restores_new_cell_text', self.read(frame, 'A8') == self.marker)
        frame.locator('.ql-quickbar [data-ql-action="bold"]').click()
        self.put(frame, 'F8', '0.125')
        self.go(frame, 'F8')
        frame.locator('.ql-quickbar [data-ql-action="more"]').click()
        frame.locator('#id-toolbar-btn-percent-style').click()
        frame.locator('.ql-back').click()
        self.go(frame, 'B4')
        sort = frame.locator('.ql-quickbar [data-ql-action="sort"]')
        self.check('quick_sort_resolves_current_native_toolbar_button', sort.is_enabled())
        sort.click()
        self.check('quick_sort_reorders_entire_filtered_range', self.read(frame, 'A4') == '墨盒' and self.read(frame, 'B4') == '1')
        frame.locator('.ql-quickbar [data-ql-action="undo"]').click()
        self.check('undo_restores_original_row_order', self.read(frame, 'A4') == '纸张')
        self.go(frame, 'B4')
        sort.click()
        self.check('sort_works_again_after_undo', self.read(frame, 'A4') == '墨盒')
        frame.locator('#status-btn-addtab').click()
        frame.wait_for_function('Asc.editor.asc_getWorksheetsCount() === 3')
        self.new_sheet = frame.evaluate('Asc.editor.asc_getWorksheetName(Asc.editor.asc_getActiveWorksheetIndex())')
        self.check('add_sheet_selects_new_native_worksheet', self.new_sheet not in ['预算明细', '汇总'])
        self.put(frame, 'A1', '新建工作表 Cross-sheet')
        self.put(frame, 'B1', "='预算明细'!D7")
        frame.locator('.ql-quickbar [data-ql-action="save"]').click()
        self.wait_saved(page, frame, self.edited, lambda book: book.sheetnames == ['预算明细', self.new_sheet, '汇总'] and book['预算明细']['B4'].value == 1 and book['预算明细']['A8'].value == self.marker)
        self.inspect_saved()
        frame.locator('#status-sheets-bar').get_by_text('预算明细', exact=True).click()
        self.style_and_motion(browser, page, frame)
        page.screenshot(path=str(self.evidence / 'spreadsheet-saved.png'))
        self.saved_hash = sha(self.edited)
        self.close()

    def reopen(self, playwright):
        old = self.pid
        self.start()
        self.check('reopen_uses_new_native_process', self.pid != old)
        self.open_argument(self.edited)
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        frame.locator('#status-sheets-bar').get_by_text('预算明细', exact=True).click()
        self.check('reopened_multi_sheet_model_and_saved_rows', frame.evaluate('Asc.editor.asc_getWorksheetsCount()') == 3 and self.read(frame, 'A4') == '墨盒' and self.read(frame, 'B6') == '5')
        self.check('reopened_text_and_formula_match_saved_file', self.read(frame, 'A8') == self.marker and self.read(frame, 'D7') == '=SUM(D4:D6)')
        self.check('fresh_reopen_preserves_file_bytes_and_clean_state', sha(self.edited) == self.saved_hash and frame.evaluate('Asc.editor.asc_isDocumentCanSave()') is False)
        page.screenshot(path=str(self.evidence / 'spreadsheet-reopened.png'))
        self.close()

    def create_new(self, playwright):
        from openpyxl import load_workbook
        self.start()
        self.open_argument('--new:cell')
        browser, page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        self.check('new_native_workbook_starts_with_one_blank_sheet', frame.evaluate('Asc.editor.asc_getWorksheetsCount()') == 1 and not self.read(frame, 'A1'))
        self.put(frame, 'A1', '新建表格验收 New workbook')
        self.put(frame, 'B1', '2')
        self.put(frame, 'C1', '3')
        self.put(frame, 'D1', '=B1+C1')
        frame.locator('.ql-quickbar [data-ql-action="save"]').click()
        dialog = owned_dialog(self.pid)
        self.check('new_workbook_uses_real_native_save_dialog', dialog.exists())
        dialog.child_window(auto_id='1001', control_type='Edit', visible_only=False).set_edit_text(str(self.created))
        dialog.child_window(auto_id='1', control_type='Button', visible_only=False).invoke()
        self.wait_saved(page, frame, self.created, lambda book: book.active['A1'].value == '新建表格验收 New workbook')
        self.check('native_save_as_keeps_uppercase_xlsx_and_adopts_path', self.created.exists() and self.created.name in [p.name for p in self.created.parent.iterdir()] and Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.created)
        book = load_workbook(self.created, data_only=True)
        try: self.check('independent_reader_confirms_new_saved_workbook_and_calculation', book.active['A1'].value == '新建表格验收 New workbook' and book.active['D1'].value == 5)
        finally: book.close()
        self.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', required=True, type=Path)
    parser.add_argument('--qt-dir', required=True, type=Path)
    parser.add_argument('--evidence-dir', required=True, type=Path)
    args = parser.parse_args()
    if os.name != 'nt': parser.error('Windows is required')
    work = args.work_dir.resolve()
    if not (work / '.qingliang-dependency-build').is_file(): raise ValueError('Select the marked stable build directory')
    build = json.loads((work / 'editor-build-report.json').read_text('utf-8'))
    if build['input_manifest_sha256'] != sha(ROOT / 'editor.inputs.lock.json'): raise ValueError('Build the current editor resources first')
    os.environ['TEMP'] = os.environ['TMP'] = str(work / 'temp')
    from playwright.sync_api import sync_playwright
    acceptance = SpreadsheetAcceptance(args)
    report = {'cases': acceptance.cases, 'observed_javascript_errors': acceptance.errors, 'normal_shutdowns': acceptance.closes,
              'style_observations': acceptance.style_observations, 'input_method': 'Chromium input events; native IME/clipboard not tested',
              'exe_sha256': sha(work / 'bin/DesktopEditors.exe'), 'host_dll_sha256': sha(work / 'bin/ascdocumentscore.dll'),
              'editor_build_report_sha256': sha(work / 'editor-build-report.json'), 'runtime_layout_sha256': sha(work / 'editor-runtime-layout.json'),
              'test_source_sha256': sha(Path(__file__)), 'native_session_helper_sha256': sha(ROOT / 'tests/run_editor_tests.py'),
              'chart_visual_acceptance': False, 'filter_value_selection_acceptance': False, 'native_clipboard_acceptance': False,
              'presentation_acceptance': False, 'pdf_forms_acceptance': False, 'installer_acceptance': False}
    try:
        acceptance.fixtures()
        with sync_playwright() as playwright:
            acceptance.edit_existing(playwright)
            acceptance.reopen(playwright)
            acceptance.create_new(playwright)
        acceptance.check('untouched_source_workbook_preserved', sha(acceptance.source) == acceptance.original_hash)
        acceptance.check('interactions_have_no_observed_javascript_errors', not acceptance.errors)
    except BaseException as error:
        report['execution_error'] = str(error)
        raise
    finally:
        if acceptance.fixtures_ready:
            report['owned_fixtures'] = {p.relative_to(acceptance.evidence).as_posix(): sha(p) for p in (acceptance.source, acceptance.edited, acceptance.created) if p.exists()}
            report['passed'] = sum(c['passed'] for c in acceptance.cases)
            report['failed'] = max(sum(not c['passed'] for c in acceptance.cases), int('execution_error' in report))
            (acceptance.evidence / 'spreadsheet-test-result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', 'utf-8')
            print(json.dumps({'passed': report['passed'], 'failed': report['failed'], 'installer_acceptance': False}), flush=True)
        elif acceptance.evidence.exists() and (acceptance.evidence / '.qingliang-spreadsheet-acceptance').exists():
            (acceptance.evidence / 'fixture-preflight-error.json').write_text(json.dumps({'execution_error': report.get('execution_error'), 'prior_ownership_report_preserved': True}, indent=2) + '\n', 'utf-8')


if __name__ == '__main__': main()
