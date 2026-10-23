"""Exercise real Qt/CEF DOCX editing, a fresh-process reopen, native Save As and PDF export.

Input uses Chromium input events; this does not claim real Windows IME candidate-window acceptance.
"""
import argparse
import ctypes
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
import urllib.request
from ctypes import wintypes as wt
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def process_kernel():
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.OpenProcess.restype = wt.HANDLE
    kernel.OpenProcess.argtypes = [wt.DWORD, wt.BOOL, wt.DWORD]
    kernel.QueryFullProcessImageNameW.argtypes = [wt.HANDLE, wt.DWORD, wt.LPWSTR, ctypes.POINTER(wt.DWORD)]
    kernel.WaitForSingleObject.argtypes = [wt.HANDLE, wt.DWORD]
    kernel.GetExitCodeProcess.argtypes = [wt.HANDLE, ctypes.POINTER(wt.DWORD)]
    kernel.CloseHandle.argtypes = [wt.HANDLE]
    return kernel


def owned_dialog(pid):
    from pywinauto import Desktop
    user = ctypes.WinDLL('user32')
    user.GetWindowThreadProcessId.argtypes = [wt.HWND, ctypes.POINTER(wt.DWORD)]
    user.GetClassNameW.argtypes = [wt.HWND, wt.LPWSTR, ctypes.c_int]
    user.IsWindowVisible.argtypes = [wt.HWND]
    result = []
    @ctypes.WINFUNCTYPE(wt.BOOL, wt.HWND, wt.LPARAM)
    def visit(hwnd, parameter):
        owner = wt.DWORD()
        user.GetWindowThreadProcessId(hwnd, ctypes.byref(owner))
        name = ctypes.create_unicode_buffer(100)
        user.GetClassNameW(hwnd, name, 100)
        if owner.value == pid and name.value == '#32770' and user.IsWindowVisible(hwnd): result.append(int(hwnd))
        return True
    for attempt in range(70):
        user.EnumWindows(visit, 0)
        if result:
            # Resolve by the verified native HWND; Qt's accessibility root can hide its owned shell dialog.
            return Desktop(backend='uia').window(handle=result[0], visible_only=False)
        time.sleep(.1)
    raise AssertionError('The owned native Save As dialog did not appear')


class Acceptance:
    def __init__(self, args):
        self.work, self.qt, self.evidence = args.work_dir.resolve(), args.qt_dir.resolve(), args.evidence_dir.resolve()
        self.cases, self.errors, self.closes = [], [], []
        self.env = dict(os.environ)
        profile = self.work / 'qt-acceptance-profile'
        self.env.update({'TEMP': str(self.work / 'temp'), 'TMP': str(self.work / 'temp'),
                         'OFFICE_USER_DATA_ROOT': str(profile / 'data'), 'QINGLIANG_SETTINGS_FILE': str(profile / 'settings.ini'),
                         'QT_QPA_PLATFORM_PLUGIN_PATH': str(self.qt / 'plugins/platforms'),
                         'VLC_PLUGIN_PATH': str(self.work / 'sources/vlc-runtime/plugins')})
        self.env['PATH'] = os.pathsep.join([str(self.work / 'bin'), str(self.work / 'bin/mediaplayer'), str(self.qt / 'bin'),
                                          str(self.work / 'sources/vlc-runtime'), self.env['PATH']])
        self.source = self.evidence / 'fixtures/source.docx'
        self.edited = self.evidence / 'fixtures/真实编辑 工作副本.docx'
        self.created = self.evidence / 'fixtures/新建保存.DOCX'
        self.pdf = self.evidence / 'fixtures/导出结果.PDF'
        self.marker = '真实桌面编辑验收 Native save 2026'
        self.new_marker = '新建并另存为验收 New document save'
        self.pid = None

    def check(self, name, value):
        self.cases.append({'name': name, 'passed': bool(value)})
        if not value: raise AssertionError(name)

    def fixtures(self):
        from docx import Document
        from docx.shared import Pt
        marker = self.evidence / '.qingliang-editor-acceptance'
        if self.evidence.exists() and any(self.evidence.iterdir()) and not marker.exists():
            raise ValueError('Select an empty evidence directory or an existing marked acceptance directory')
        self.evidence.mkdir(parents=True, exist_ok=True)
        marker.write_text('Only controlled editor acceptance fixtures and evidence.\n', 'utf-8')
        (self.evidence / 'fixtures').mkdir(exist_ok=True)
        result = self.evidence / 'editor-test-result.json'
        prior = json.loads(result.read_text('utf-8')) if result.exists() else {}
        known = prior.get('owned_fixtures', {})
        for path in (self.source, self.edited, self.created, self.pdf):
            if path.exists():
                if not path.resolve().is_relative_to(self.evidence) or sha(path) != known.get(path.relative_to(self.evidence).as_posix()):
                    raise ValueError('Unknown existing fixture edit preserved: ' + path.name)
                if path != self.source: path.unlink()
        if not self.source.exists():
            doc = Document()
            doc.core_properties.author = 'Qingliang acceptance'
            doc.styles['Normal'].font.name = 'Calibri'
            doc.styles['Normal'].font.size = Pt(12)
            doc.add_paragraph('Original desktop document.')
            doc.add_paragraph('原有中文段落。')
            doc.save(self.source)
        shutil.copy2(self.source, self.edited)
        self.original_hash = sha(self.source)

    def start(self):
        subprocess.run([sys.executable, str(ROOT / 'tools/launch_qt_preview.py'), '--work-dir', str(self.work),
                        '--qt-dir', str(self.qt), '--gui'], cwd=ROOT, env=self.env, check=True,
                       stdout=subprocess.DEVNULL)
        launch = json.loads((self.work / 'qt-runtime-launch.json').read_text('utf-8'))
        self.pid, self.port = launch['pid'], launch['debug_port']

    def close(self):
        kernel, user = process_kernel(), ctypes.WinDLL('user32')
        user.GetWindowThreadProcessId.argtypes = [wt.HWND, ctypes.POINTER(wt.DWORD)]
        user.PostMessageW.argtypes = [wt.HWND, wt.UINT, wt.WPARAM, wt.LPARAM]
        user.IsWindowVisible.argtypes = [wt.HWND]
        user.GetWindowTextLengthW.argtypes = [wt.HWND]
        handle = kernel.OpenProcess(0x101000, False, self.pid)
        if not handle: raise ctypes.WinError(ctypes.get_last_error())
        try:
            image, size = ctypes.create_unicode_buffer(32768), wt.DWORD(32768)
            if not kernel.QueryFullProcessImageNameW(handle, 0, image, ctypes.byref(size)) or Path(image.value).resolve() != self.work / 'bin/DesktopEditors.exe':
                raise ValueError('The recorded acceptance process is not the staged desktop')
            windows = []
            @ctypes.WINFUNCTYPE(wt.BOOL, wt.HWND, wt.LPARAM)
            def visit(hwnd, parameter):
                owner = wt.DWORD()
                user.GetWindowThreadProcessId(hwnd, ctypes.byref(owner))
                if owner.value == self.pid and user.IsWindowVisible(hwnd) and user.GetWindowTextLengthW(hwnd): windows.append(hwnd)
                return True
            user.EnumWindows(visit, 0)
            if not windows: raise AssertionError('No owned preview window to close')
            for hwnd in windows: user.PostMessageW(hwnd, 0x0010, 0, 0)
            if kernel.WaitForSingleObject(handle, 20000) != 0:
                raise AssertionError('Preview did not close normally; it has not been force-terminated')
            code = wt.DWORD()
            if not kernel.GetExitCodeProcess(handle, ctypes.byref(code)): raise ctypes.WinError(ctypes.get_last_error())
            self.closes.append({'pid': self.pid, 'exit_code': code.value, 'force_termination_used': False})
            self.check('normal_native_shutdown_' + str(len(self.closes)), code.value == 0)
        finally: kernel.CloseHandle(handle)

    def open_argument(self, argument):
        result = subprocess.run([str(self.work / 'bin/DesktopEditors.exe'), '--log', str(argument)],
                                cwd=self.work / 'bin', env=self.env, capture_output=True, timeout=15)
        self.check('native_open_command_' + str(len(self.closes)), result.returncode == 0)
        for attempt in range(50):
            with urllib.request.urlopen('http://127.0.0.1:' + str(self.port) + '/json/list', timeout=3) as response: targets = json.load(response)
            if any('/apps/api/documents/' in row.get('url', '') for row in targets): return
            time.sleep(.1)
        raise AssertionError('The native editor browser target did not appear')

    def connect(self, playwright):
        browser = playwright.chromium.connect_over_cdp('http://127.0.0.1:' + str(self.port))
        page = next(p for context in browser.contexts for p in context.pages if '/apps/api/documents/' in p.url)
        page.set_default_timeout(7000)
        page.on('pageerror', lambda error: self.errors.append(str(error)))
        frame = None
        for attempt in range(80):
            frame = next((f for f in page.frames if '/documenteditor/main/' in f.url), None)
            if frame and frame.evaluate("(() => {try {return typeof window.editor?.asc_nativePrintPagesCount === 'function' && editor.asc_nativePrintPagesCount() > 0;} catch {return false;}})()"): return page, frame
            page.wait_for_timeout(250)
        raise AssertionError('The real Word editor did not finish rendering')

    def dismiss_tips(self, frame):
        tips = frame.locator('.synch-tip-root:visible')
        for attempt in range(tips.count()):
            button = tips.first.get_by_text('知道了', exact=True)
            if button.count(): button.click()

    def text(self, page, frame):
        frame.locator('#id_viewer_overlay').click(position={'x': 350, 'y': 130})
        page.keyboard.press('Control+A')
        return frame.evaluate('editor.pluginMethod_GetSelectedText({Numbering:false,ParaSeparator:"\\n"})')

    def edit_existing(self, playwright):
        from docx import Document
        self.start()
        self.open_argument(self.edited)
        page, frame = self.connect(playwright)
        self.check('existing_docx_native_source_path', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.edited)
        self.dismiss_tips(frame)
        initial = self.text(page, frame)
        self.check('existing_native_model_contains_original_text', 'Original desktop document.' in initial and '原有中文段落。' in initial)
        before = sha(self.edited)
        page.keyboard.press('Control+End')
        page.keyboard.press('Enter')
        page.keyboard.insert_text(self.marker)
        frame.wait_for_function('editor.asc_isDocumentCanSave() === true')
        self.check('editing_enables_save', frame.locator('.ql-quickbar [data-ql-action="save"]').is_enabled())
        page.keyboard.press('Control+Z')
        self.check('undo_removes_inserted_text', self.marker not in self.text(page, frame))
        frame.locator('.ql-quickbar [data-ql-action="redo"]').click()
        self.check('redo_restores_inserted_text', self.marker in self.text(page, frame))
        page.keyboard.press('Control+End')
        page.keyboard.press('Home')
        page.keyboard.press('Shift+End')
        self.check('keyboard_selects_only_inserted_paragraph', frame.evaluate('editor.pluginMethod_GetSelectedText({Numbering:false})').strip() == self.marker)
        frame.locator('.ql-quickbar [data-ql-action="bold"]').click()
        self.check('bold_action_updates_state', frame.locator('.ql-quickbar [data-ql-action="bold"]').get_attribute('aria-pressed') == 'true')
        frame.locator('.ql-quickbar [data-ql-action="save"]').click()
        saved = None
        for attempt in range(100):
            page.wait_for_timeout(150)
            try:
                if sha(self.edited) != before:
                    candidate = Document(self.edited)
                    if any(self.marker in p.text for p in candidate.paragraphs): saved = candidate; break
            except (OSError, ValueError, KeyError): pass
        self.check('native_save_writes_readable_docx', saved is not None)
        self.check('save_preserves_original_paragraphs', [p.text for p in saved.paragraphs][:2] == ['Original desktop document.', '原有中文段落。'])
        paragraph = next(p for p in saved.paragraphs if self.marker in p.text)
        self.check('independent_reader_confirms_bold', all(run.bold is True for run in paragraph.runs if run.text.strip()))
        frame.wait_for_function('editor.asc_isDocumentCanSave() === false')
        self.check('save_clears_dirty_state', frame.locator('.ql-quickbar [data-ql-action="save"]').is_disabled())
        page.screenshot(path=str(self.evidence / 'word-saved.png'))
        self.saved_hash = sha(self.edited)
        self.close()

    def reopen(self, playwright):
        old = self.pid
        self.start()
        self.check('reopen_uses_a_new_native_process', self.pid != old)
        self.open_argument(self.edited)
        page, frame = self.connect(playwright)
        self.dismiss_tips(frame)
        text = self.text(page, frame)
        self.check('fresh_native_model_contains_saved_text', self.marker in text and 'Original desktop document.' in text and '原有中文段落。' in text)
        page.keyboard.press('Control+End')
        page.keyboard.press('Home')
        page.keyboard.press('Shift+End')
        self.check('reopened_bold_reflected_in_toolbar', frame.locator('.ql-quickbar [data-ql-action="bold"]').get_attribute('aria-pressed') == 'true')
        self.check('reopen_preserves_file_bytes', sha(self.edited) == self.saved_hash)
        self.check('reopened_document_is_clean', frame.evaluate('editor.asc_isDocumentCanSave()') is False)
        page.screenshot(path=str(self.evidence / 'word-reopened.png'))
        self.close()

    def create_and_export(self, playwright):
        from docx import Document
        import fitz
        self.start()
        self.open_argument('--new:word')
        page, frame = self.connect(playwright)
        tip = frame.locator('.synch-tip-root:visible')
        if not tip.count():
            frame.evaluate("(() => {const t=DE.getController('Toolbar');Common.UI.TooltipManager.showTip({step:'qa-multipage-hint',target:'.ql-quickbar',placement:'bottom-left',text:t.helpMultipageView,header:t.helpMultipageViewHeader,isNewFeature:true,showButton:true,noArrow:true,multiple:true,closable:false});})()")
        text = tip.first.inner_text()
        self.check('localized_help_has_correct_chinese_heading', '新功能' in text and '多页视图' in text)
        self.check('localized_help_has_correct_feature_content', '同时查看多个页面' in text)
        self.check('help_has_no_undefined_or_unrelated_collaboration_warning', 'undefined' not in text and '文档已被其他用户更改' not in text)
        page.screenshot(path=str(self.evidence / 'word-help.png'))
        self.dismiss_tips(frame)
        self.check('native_new_template_is_empty', not self.text(page, frame).strip())
        page.keyboard.insert_text(self.new_marker)
        frame.wait_for_function('editor.asc_isDocumentCanSave() === true')
        frame.locator('.ql-quickbar [data-ql-action="save"]').click()
        dialog = owned_dialog(self.pid)
        self.check('new_save_opens_real_native_file_dialog', dialog.exists())
        dialog.child_window(auto_id='1001', control_type='Edit', visible_only=False).set_edit_text(str(self.created))
        dialog.child_window(auto_id='1', control_type='Button', visible_only=False).invoke()
        saved = False
        for attempt in range(100):
            page.wait_for_timeout(150)
            if self.created.exists():
                try:
                    saved = self.new_marker in '\n'.join(p.text for p in Document(self.created).paragraphs)
                    if saved: break
                except (OSError, ValueError, KeyError): pass
        self.check('native_save_as_writes_real_docx', saved)
        self.check('save_as_keeps_uppercase_extension', self.created.name in [p.name for p in self.created.parent.iterdir()])
        self.check('native_source_adopts_save_as_path', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.created)
        frame.wait_for_function('editor.asc_isDocumentCanSave() === false')
        self.check('new_save_clears_dirty_state', frame.locator('.ql-quickbar [data-ql-action="save"]').is_disabled())
        page.screenshot(path=str(self.evidence / 'word-created.png'))
        docx_hash = sha(self.created)
        frame.locator('.ql-quickbar [data-ql-action="file"]').click()
        frame.locator('.ql-command-item[data-ql-action="exportpdf"]').click()
        dialog = owned_dialog(self.pid)
        self.check('pdf_export_opens_real_native_file_dialog', dialog.exists())
        dialog.child_window(auto_id='1001', control_type='Edit', visible_only=False).set_edit_text(str(self.pdf))
        dialog.child_window(auto_id='1', control_type='Button', visible_only=False).invoke()
        pdf_text = ''
        for attempt in range(150):
            page.wait_for_timeout(150)
            if self.pdf.exists():
                try:
                    with fitz.open(self.pdf) as document:
                        pdf_text = '\n'.join(p.get_text() for p in document)
                        if pdf_text.strip():
                            document[0].get_pixmap(matrix=fitz.Matrix(1.2, 1.2)).save(str(self.evidence / 'exported-pdf-page.png'))
                            break
                except (OSError, ValueError, RuntimeError): pass
        self.check('native_pdf_is_independently_readable', bool(pdf_text.strip()))
        self.check('pdf_contains_actual_chinese_and_english', '新建并另存为验收' in pdf_text and 'New document save' in pdf_text)
        self.check('pdf_keeps_uppercase_extension', self.pdf.name in [p.name for p in self.pdf.parent.iterdir()])
        self.check('pdf_export_preserves_original_docx', sha(self.created) == docx_hash)
        self.check('pdf_export_keeps_docx_editing_context', Path(frame.evaluate('AscDesktopEditor.LocalFileGetSourcePath()')).resolve() == self.created)
        self.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--work-dir', required=True, type=Path)
    parser.add_argument('--qt-dir', required=True, type=Path)
    parser.add_argument('--evidence-dir', required=True, type=Path)
    args = parser.parse_args()
    if os.name != 'nt': parser.error('Windows is required')
    work = args.work_dir.resolve()
    if not (work / '.qingliang-dependency-build').is_file(): raise ValueError('Select the marked stable editor build directory')
    build = json.loads((work / 'editor-build-report.json').read_text('utf-8'))
    if build['input_manifest_sha256'] != sha(ROOT / 'editor.inputs.lock.json'): raise ValueError('Build the current desktop editor resources first')
    os.environ['TEMP'] = os.environ['TMP'] = str(work / 'temp')
    os.environ['PIP_CACHE_DIR'] = str(work / 'pip-cache')
    from playwright.sync_api import sync_playwright
    acceptance = Acceptance(args)
    report = {'cases': acceptance.cases, 'observed_javascript_errors': acceptance.errors, 'normal_shutdowns': acceptance.closes,
              'input_method': 'Chromium input events; Windows IME candidate window not tested',
              'exe_sha256': sha(work / 'bin/DesktopEditors.exe'), 'host_dll_sha256': sha(work / 'bin/ascdocumentscore.dll'),
              'editor_build_report_sha256': sha(work / 'editor-build-report.json'), 'runtime_layout_sha256': sha(work / 'editor-runtime-layout.json'),
              'test_source_sha256': sha(Path(__file__)), 'spreadsheet_acceptance': False, 'presentation_acceptance': False,
              'pdf_form_editing_acceptance': False, 'installer_acceptance': False}
    try:
        acceptance.fixtures()
        with sync_playwright() as playwright:
            acceptance.edit_existing(playwright)
            acceptance.reopen(playwright)
            acceptance.create_and_export(playwright)
        acceptance.check('untouched_source_docx_preserved', sha(acceptance.source) == acceptance.original_hash)
        acceptance.check('interactions_have_no_observed_javascript_errors', not acceptance.errors)
    except BaseException as error:
        report['execution_error'] = str(error)
        raise
    finally:
        if acceptance.evidence.exists() and (acceptance.evidence / '.qingliang-editor-acceptance').exists():
            report['owned_fixtures'] = {p.relative_to(acceptance.evidence).as_posix(): sha(p) for p in (acceptance.source, acceptance.edited, acceptance.created, acceptance.pdf) if p.exists()}
            report['passed'] = sum(c['passed'] for c in acceptance.cases)
            report['failed'] = max(sum(not c['passed'] for c in acceptance.cases), int('execution_error' in report))
            (acceptance.evidence / 'editor-test-result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', 'utf-8')
            print(json.dumps({'passed': report['passed'], 'failed': report['failed'], 'installer_acceptance': False}), flush=True)


if __name__ == '__main__':
    main()
