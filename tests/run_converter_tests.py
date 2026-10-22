"""Exercise the actual x2t executable and read outputs with independent libraries."""
import argparse
import csv
import io
import json
import os
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path
import docx
import openpyxl
import pptx
from pptx.util import Inches
import pymupdf


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--exe',type=Path,required=True);parser.add_argument('--work-dir',type=Path,required=True)
    parser.add_argument('--reuse-fixtures',action='store_true',help='Explicitly reuse this generated fixture directory')
    args=parser.parse_args();exe=args.exe.resolve();folder=args.work_dir.resolve()/'converter-fixtures'
    folder.mkdir(exist_ok=True)
    if any(folder.iterdir()) and not args.reuse_fixtures:raise ValueError('Prior converter evidence remains; inspect before reusing')
    details=[];counter=0;fonts=exe.parent/'sdkjs/common';tmp=folder/'temp';tmp.mkdir(exist_ok=True)
    env=os.environ.copy();env.update({'TEMP':str(tmp),'TMP':str(tmp)})
    def check(condition,message):
        if not condition:raise AssertionError(message)
    def convert(source,output,encoding='46',expect_ok=True):
        nonlocal counter
        counter+=1;scratch=tmp/str(counter);scratch.mkdir(exist_ok=True);xml=ET.Element('TaskQueueDataConvert')
        values={'m_sFileFrom':str(source),'m_sFileTo':str(output),'m_sTempDir':str(scratch),'m_nCsvTxtEncoding':encoding,'m_nCsvDelimiter':'4','m_sDefaultFontName':'Noto Sans CJK SC','m_sFontDir':str(fonts),'m_sAllFontsPath':str(fonts/'AllFonts.js')}
        for key,value in values.items():ET.SubElement(xml,key).text=value
        params=folder/('params-'+str(counter)+'.xml');ET.ElementTree(xml).write(params,encoding='utf-8',xml_declaration=True)
        result=subprocess.run([str(exe),str(params)],cwd=exe.parent,env=env,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=180)
        (folder/('process-'+str(counter)+'.log')).write_bytes(result.stdout)
        check((result.returncode==0)==expect_ok,'x2t exit '+str(result.returncode)+' for '+source.name+' -> '+output.name)
        if expect_ok:check(output.is_file() and output.stat().st_size>0,'No output: '+output.name)
        return result
    def case(name,action):
        try:action();row={'name':name,'passed':True}
        except Exception as error:row={'name':name,'passed':False,'error':repr(error)}
        details.append(row);print(('PASS ' if row['passed'] else 'FORMAT_FAILURE ')+json.dumps(row,ensure_ascii=True),flush=True)
    def doc_text(path):return '\n'.join(p.text for p in docx.Document(path).paragraphs)
    def pdf_text(path):
        with pymupdf.open(path) as pdf:
            check(len(pdf)>0,'Empty PDF');text=''.join(p.get_text() for p in pdf)
            check(not pdf.is_repaired,'PDF required repairs')
            pix=pdf[0].get_pixmap(matrix=pymupdf.Matrix(1,1),alpha=False)
            check(min(pix.samples)<150,'PDF page is blank');pix.save(str(path.with_suffix('.png')))
            return text
    samples=[('ascii','Hello x2t'),('unicode','中文 日本語 Ελληνικά 😀'),('xml-special','A & B <tag> "quote" > end'),
             ('tabs','\tA\t\tB\t'),('empty-lines','first\n\nthird'),('mixed-newlines','first\r\nsecond\rthird\nlast'),
             ('spaces','  leading and trailing  '),('bom','\ufeff带 BOM 文本'),('trailing-newline','one\n')]
    for name,text in samples:
        def roundtrip(name=name,text=text):
            source=folder/(name+' 中文.txt');target=folder/(name+'.docx');back=folder/(name+'-back.txt');source.write_bytes(text.encode('utf-8'))
            convert(source,target);normalized=text.removeprefix('\ufeff').replace('\r\n','\n').replace('\r','\n').removesuffix('\n')
            check(doc_text(target)==normalized,'Independent DOCX text differs');convert(target,back)
            check(back.read_bytes().decode('utf-8-sig').replace('\r\n','\n')==normalized,'TXT round trip differs')
        case('TXT/DOCX '+name,roundtrip)
    independent=folder/'independent.docx';d=docx.Document();d.add_paragraph('Converter acceptance 2026');d.add_paragraph('中文转换验证');table=d.add_table(rows=1,cols=2);table.cell(0,0).text='Left';table.cell(0,1).text='Right';d.save(independent)
    d.styles['Normal'].font.name='Noto Sans CJK SC';d.styles['Normal']._element.rPr.rFonts.set(docx.oxml.ns.qn('w:eastAsia'),'Noto Sans CJK SC');d.save(independent)
    def table_test():
        target=folder/'table.txt';convert(independent,target);text=target.read_bytes().decode('utf-8-sig')
        check('Converter acceptance 2026' in text and '中文转换验证' in text and '|Left |Right |' in text,'Table extraction differs')
    case('Independent DOCX paragraphs and table to TXT',table_test)
    def word_pdf():
        target=folder/'word.pdf';convert(independent,target);text=pdf_text(target)
        check('Converter acceptance 2026' in text and '中文转换验证' in text and 'Left' in text and 'Right' in text,'PDF text differs')
    case('Independent DOCX to PDF with Chinese text and table',word_pdf)
    for ext in ['rtf','odt']:
        def roundtrip_format(ext=ext):
            target=folder/('word.'+ext);back=folder/('word-'+ext+'-back.docx');convert(independent,target);convert(target,back);d=docx.Document(back)
            check('Converter acceptance 2026' in doc_text(back) and '中文转换验证' in doc_text(back),'Paragraphs lost in '+ext)
            check(len(d.tables)==1 and d.tables[0].cell(0,1).text=='Right','Table lost in '+ext)
        case('DOCX/'+ext.upper()+' actual round trip',roundtrip_format)
    workbook=folder/'independent.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='Test';sheet.append(['Item','Amount']);sheet.append(['Alpha',123.5]);sheet.append(['中文',7]);book.save(workbook)
    for row in sheet:
        for cell in row:cell.font=openpyxl.styles.Font(name='Noto Sans CJK SC',size=11)
    book.save(workbook)
    def spreadsheet_csv():
        target=folder/'sheet.csv';convert(workbook,target);rows=list(csv.reader(io.StringIO(target.read_bytes().decode('utf-8-sig'))))
        check(rows==[['Item','Amount'],['Alpha','123.5'],['中文','7']],'CSV cells differ: '+repr(rows))
    case('Independent XLSX to UTF8 CSV',spreadsheet_csv)
    def csv_xlsx():
        source=folder/'input.csv';source.write_bytes('Name,Value\n"中文,quoted",42\n'.encode('utf-8'));target=folder/'csv.xlsx';convert(source,target)
        sheet=openpyxl.load_workbook(target).active
        check(sheet['A2'].value=='中文,quoted' and sheet['B2'].value==42,'CSV import cells differ')
    case('UTF8 CSV quoted cell to independent XLSX reader',csv_xlsx)
    def spreadsheet_ods():
        target=folder/'sheet.ods';back=folder/'sheet-ods-back.xlsx';convert(workbook,target);convert(target,back);sheet=openpyxl.load_workbook(back).active
        check(list(sheet.values)[:3]==[('Item','Amount'),('Alpha',123.5),('中文',7)],'ODS round trip cells differ: '+repr(list(sheet.values)[:3]))
    case('XLSX/ODS actual round trip',spreadsheet_ods)
    def spreadsheet_pdf():
        target=folder/'sheet.pdf';convert(workbook,target);text=pdf_text(target)
        check('Alpha' in text and '123.5' in text and '中文' in text,'Spreadsheet PDF text differs')
    case('Independent XLSX to PDF',spreadsheet_pdf)
    presentation=folder/'independent.pptx';slides=pptx.Presentation();slide=slides.slides.add_slide(slides.slide_layouts[6]);shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2));shape.text_frame.text='Presentation acceptance';shape.text_frame.add_paragraph().text='中文演示验证';slides.save(presentation)
    for paragraph in shape.text_frame.paragraphs:paragraph.font.name='Noto Sans CJK SC'
    slides.save(presentation)
    def presentation_pdf():
        target=folder/'slide.pdf';convert(presentation,target);text=pdf_text(target)
        check('Presentation acceptance' in text and '中文演示验证' in text,'Presentation PDF text differs')
    case('Independent PPTX to PDF',presentation_pdf)
    def presentation_odp():
        target=folder/'slide.odp';back=folder/'slide-odp-back.pptx';convert(presentation,target);convert(target,back)
        text='\n'.join(s.text for slide in pptx.Presentation(back).slides for s in slide.shapes if s.has_text_frame)
        check('Presentation acceptance' in text and '中文演示验证' in text,'ODP round trip text differs')
    case('PPTX/ODP actual round trip',presentation_odp)
    missing=folder/'missing.docx';bad=folder/'truncated-zip.docx';bad.write_bytes(independent.read_bytes()[:100])
    for name,source in [('missing',missing),('corrupt',bad)]:
        def reject(name=name,source=source):
            output=folder/(source.stem+'-rejected.txt');convert(source,output,expect_ok=False);check(not output.exists(),'Invalid input created output')
        case('Reject '+name+' input without output',reject)
    for value in ['-1','999999999999999999999999','46junk','1000']:
        def reject_encoding(value=value):
            source=folder/'invalid-encoding.txt';source.write_bytes(b'original');output=folder/('encoding-'+str(len(details))+'.docx');output.write_bytes(b'KEEP ORIGINAL');convert(source,output,encoding=value,expect_ok=False)
            check(output.read_bytes()==b'KEEP ORIGINAL','Invalid encoding overwrote output')
        case('Reject encoding '+value+' and preserve existing output',reject_encoding)
    result={'passed':sum(r['passed'] for r in details),'failed':sum(not r['passed'] for r in details),'conversion_processes':counter,'exe_sha256':__import__('hashlib').sha256(exe.read_bytes()).hexdigest()}
    (args.work_dir/'converter-native-cases.json').write_text(json.dumps(details,indent=2)+'\n','utf-8');(args.work_dir/'converter-test-result.json').write_text(json.dumps(result,indent=2)+'\n','utf-8');print(json.dumps(result),flush=True)
    if result['failed']:raise RuntimeError('Converter acceptance failed; fixtures and logs retained')


if __name__=='__main__':main()
